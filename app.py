"""
Thalassemia WebApp - Image Acquisition App
Jalankan:
    pip install -r requirements.txt
    python app.py
Lalu buka http://127.0.0.1:5000 di browser.
"""

import base64
import csv
import json
import math
import os
import platform
import random
import re
import sys
import threading
import time
from datetime import datetime

import cv2
import numpy as np
from flask import Flask, Response, jsonify, render_template, request

app = Flask(__name__)

# Untuk deteksi buram
FOCUS_BLUR_THRESHOLD = 60.0

# untuk deteksi objek tidak ada
BLANK_STD_THRESHOLD = 5.0

# Berapa kali read() OpenCV boleh gagal berturut-turut sebelum dianggap
OPENCV_DISCONNECT_THRESHOLD = 15
LOG_FILENAME = "capture_log.csv"

# Lindungi critical section /api/save (hitung iterasi + tulis file + tulis
# log CSV) dari race condition kalau ada dua request /api/save nyaris
# bersamaan (Flask jalan threaded=True) -- tanpa ini, dua request bisa saja
# baca "iterasi berikutnya" yang sama sebelum salah satu selesai menulis,
# jadi filename ketimpa/duplikat atau baris CSV keselip.
save_lock = threading.Lock()

# Label kelas morfologi sel eritrosit untuk tab Segmentasi & Klasifikasi.
# Urutan index di sini PENTING -- ini bakal jadi mapping output model
# klasifikasi (deep learning) begitu pipeline-nya dipasang, jadi jangan
# diubah urutannya tanpa nyesuaikan juga di sisi model/training nanti.
RBC_CLASS_LABELS = [
    "Normal cell",
    "Macrocyte",
    "Microcyte",
    "Spherocyte",
    "Target cell",
    "Stomatocyte",
    "Ovalocyte",
    "Teardrop",
    "Burr cell",
    "Schistocyte",
    "Uncategorised",
    "Hypochromia",
    "Elliptocyte",
]

# Warna tampilan per kelas (dipakai buat overlay & legenda di frontend) --
# ditaruh satu sumber di backend biar konsisten kalau nanti dipakai juga di
# skrip training/evaluasi (mis. visualisasi confusion matrix per warna kelas
# yang sama).
RBC_CLASS_COLORS = [
    "#16a34a",  # Normal cell
    "#2563eb",  # Macrocyte
    "#f59e0b",  # Microcyte
    "#db2777",  # Spherocyte
    "#7c3aed",  # Target cell
    "#0891b2",  # Stomatocyte
    "#ea580c",  # Ovalocyte
    "#65a30d",  # Teardrop
    "#dc2626",  # Burr cell
    "#9333ea",  # Schistocyte
    "#64748b",  # Uncategorised
    "#ca8a04",  # Hypochromia
    "#0d9488",  # Elliptocyte
]

# File ringkasan hasil segmentasi (satu baris per proses "Simpan Hasil"),
# ditulis di folder yang sama dengan gambar sumbernya -- polanya sengaja
# dibuat mirip capture_log.csv biar gampang dianalisis bareng.
SEGMENTATION_LOG_FILENAME = "segmentation_log.csv"

# is_duplicate/duplicate_of: ditambahkan belakangan (lihat _ensure_log_columns)
# buat nandain kalau suatu hasil disimpan persis sama (jumlah sel per kelas +
# total identik) dengan hasil TERAKHIR yang sudah tersimpan buat gambar yang
# sama -- biar folder nggak numpuk file JSON duplikat tanpa keterangan.
#
# doctor_reviewed_count/doctor_correct_count/doctor_incorrect_count: agregasi
# penilaian dokter per sel (lihat doctorVerdict di detections) -- diisi 0 buat
# hasil batch (mock otomatis, tidak pernah direview satu-satu).
SEGMENTATION_LOG_FIELDNAMES = (
    ["timestamp", "patient_id", "source_image", "detail_file", "total_cells", "is_duplicate", "duplicate_of"]
    + ["doctor_reviewed_count", "doctor_correct_count", "doctor_incorrect_count"]
    + RBC_CLASS_LABELS
)

# Lindungi penulisan file detail JSON + baris segmentation_log.csv dari race
# condition kalau ada beberapa "Simpan Hasil" nyaris bersamaan -- pola yang
# sama dengan save_lock di atas.
segmentation_log_lock = threading.Lock()


def _ensure_log_columns(log_path, fieldnames):
    """Migrasi in-place kalau segmentation_log.csv yang sudah ada (dari versi
    lama) belum punya kolom baru (mis. is_duplicate/duplicate_of) -- baca
    semua baris lama, rewrite dengan header terbaru (kolom baru diisi kosong
    utk baris lama), biar CSV tetap satu format konsisten & nggak ada nilai
    ke-shift/hilang pas dibuka di Excel setelah kolom baru ditambahkan.
    Harus dipanggil di dalam segmentation_log_lock."""
    if not os.path.isfile(log_path):
        return
    with open(log_path, "r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        header = reader.fieldnames or []
        rows = list(reader)
    if not header or all(name in header for name in fieldnames):
        return  # file kosong, atau sudah punya semua kolom yg dibutuhkan
    with open(log_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            row.pop(None, None)  # buang overflow restkey (baris lama yg kolomnya beda)
            writer.writerow({name: row.get(name, "") for name in fieldnames})


# ---------------------------------------------------------------------------
# SDK vendor ToupCam (opsional, cuma dipakai di Linux)
# ---------------------------------------------------------------------------
# Beberapa kamera vendor (mis. MiiCam, rebrand dari chipset ToupTek) ternyata
# tidak sepenuhnya UVC-compliant dari sisi yang bisa di-bind modul kernel
# `uvcvideo` standar -- di beberapa board (mis. NVIDIA Jetson) device-nya
# kebaca di level USB tapi gagal dapat /dev/videoN sama sekali. SDK vendor
# (komunikasi langsung ke USB lewat libusb, bypass uvcvideo/V4L2) berhasil
# baca kameranya. Ini loader lazy buat SDK itu.
#
# libtoupcam.so adalah biner native per-arsitektur CPU, jadi dibundel
# per-arsitektur di vendor/toupcam/<arch>/ (masing-masing juga bawa salinan
# toupcam.py sendiri, karena binding-nya cari .so di folder yang sama
# persis dengan dirinya). Arsitektur dideteksi otomatis lewat
# platform.machine() -- tinggal tambah folder arch baru kalau nanti perlu
# platform lain (mis. armhf untuk Raspberry Pi 32-bit).
#
# Cuma dicoba di Linux -- di Windows OpenCV+DirectShow sudah cukup, jadi
# tidak diutak-atik supaya tidak menambah risiko regresi di sana.
_toupcam_sdk_cache = {"loaded": False, "module": None}

_TOUPCAM_ARCH_MAP = {
    "aarch64": "arm64",
    "arm64": "arm64",
    "x86_64": "x64",
    "amd64": "x64",
}


def _load_toupcam_sdk():
    if _toupcam_sdk_cache["loaded"]:
        return _toupcam_sdk_cache["module"]
    _toupcam_sdk_cache["loaded"] = True
    if os.name != "posix":
        return None

    arch_dir = _TOUPCAM_ARCH_MAP.get(platform.machine().lower())
    if arch_dir is None:
        print(f"[DEBUG] ToupCam SDK: arsitektur '{platform.machine()}' belum didukung/dibundel.")
        return None

    vendor_dir = os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "vendor", "toupcam", arch_dir
    )
    if not os.path.isfile(os.path.join(vendor_dir, "libtoupcam.so")):
        return None
    try:
        if vendor_dir not in sys.path:
            sys.path.insert(0, vendor_dir)
        import toupcam as toupcam_sdk

        _toupcam_sdk_cache["module"] = toupcam_sdk
        return toupcam_sdk
    except Exception as e:
        print(f"[DEBUG] ToupCam SDK gagal dimuat: {e}")
        return None


# ---------------------------------------------------------------------------
# Camera state
# ---------------------------------------------------------------------------
class CameraManager:
    def __init__(self):
        self.lock = threading.Lock()
        self.backend = None          # "opencv" atau "toupcam", None kalau belum connect
        self.cap = None              # cv2.VideoCapture, dipakai kalau backend == "opencv"
        self.current_index = None    # index tampilan (posisi di list_cameras())
        self.current_width = None    # resolusi aktual kamera yang sedang terhubung
        self.current_height = None
        self.last_frame = None       # frame terakhir dari stream (untuk MJPEG)
        self.captured_frame = None   # frame yang dibekukan lewat tombol Capture
        self.captured_focus_score = None
        self._registry = []          # hasil list_cameras() terakhir, index -> {backend, ...}

        # State khusus backend ToupCam
        self._tc_handle = None
        self._tc_buf = None
        self._tc_row_pitch = 0
        self._tc_width = 0
        self._tc_height = 0
        self._tc_connected = False
        self._tc_last_pull = 0.0

        # State khusus backend OpenCV -- cv2.VideoCapture.isOpened() TIDAK
        # otomatis jadi False cuma karena kameranya dicabut fisik; harus
        # dideteksi lewat read() yang gagal berkali-kali berturut-turut.
        self._opencv_connected = False
        self._opencv_read_failures = 0

    def list_cameras(self, max_index=5):
        """Daftar kamera yang tersedia, gabungan SDK vendor ToupCam (kalau ada
        kamera yang cuma bisa dibaca lewat situ, mis. MiiCam di Linux) dan
        probe generik OpenCV/pygrabber/V4L2. Hasilnya disimpan di
        self._registry supaya open(index) tahu backend & device id yang tepat
        untuk index yang sama persis dengan yang ditampilkan di sini."""
        registry = []

        toupcam_sdk = _load_toupcam_sdk()
        if toupcam_sdk is not None:
            try:
                devices = toupcam_sdk.Toupcam.EnumV2()
                for dev in devices:
                    registry.append({
                        "name": dev.displayname,
                        "backend": "toupcam",
                        "toupcam_id": dev.id,
                    })
                if devices:
                    print(f"[DEBUG] ToupCam SDK menemukan device: {[d.displayname for d in devices]}")
            except Exception as e:
                print(f"[DEBUG] ToupCam EnumV2 gagal: {e}")

        if os.name == "nt":
            try:
                import comtypes
                try:
                    comtypes.CoInitialize()
                except Exception:
                    pass

                from pygrabber.dshow_graph import FilterGraph

                names = FilterGraph().get_input_devices()
                print(f"[DEBUG] pygrabber menemukan device: {list(enumerate(names))}")
                for i, n in enumerate(names):
                    registry.append({"name": n, "backend": "opencv", "cv_index": i})
            except Exception as e:
                print(f"[DEBUG] pygrabber gagal, fallback ke probe OpenCV: {e}")
                registry.extend(self._probe_opencv_indices(max_index))
        else:
            registry.extend(self._probe_opencv_indices(max_index))

        for i, entry in enumerate(registry):
            entry["index"] = i
        self._registry = registry
        return [{"index": e["index"], "name": e["name"]} for e in registry]

    def _probe_opencv_indices(self, max_index):
        backend = cv2.CAP_DSHOW if os.name == "nt" else 0
        found = []
        for i in range(max_index):
            cap = cv2.VideoCapture(i, backend)
            if cap is not None and cap.isOpened():
                name = f"Kamera USB #{i}"
                if os.name == "posix":
                    name_path = f"/sys/class/video4linux/video{i}/name"
                    try:
                        if os.path.isfile(name_path):
                            with open(name_path, "r") as f:
                                real_name = f.read().strip()
                            if real_name:
                                name = real_name
                    except OSError:
                        pass
                found.append({"name": name, "backend": "opencv", "cv_index": i})
            if cap is not None:
                cap.release()
        return found

    def open(self, index, width=None, height=None):
        """Buka kamera pada index tampilan (hasil list_cameras()). Otomatis
        pilih backend yang sesuai (ToupCam SDK atau OpenCV) berdasarkan
        self._registry."""
        with self.lock:
            entry = self._registry[index] if 0 <= index < len(self._registry) else None
            self._release_locked()

            if entry is not None and entry["backend"] == "toupcam":
                result = self._open_toupcam_locked(entry["toupcam_id"], width, height)
            else:
                cv_index = entry["cv_index"] if entry is not None else index
                result = self._open_opencv_locked(cv_index, width, height)

            self.current_index = index if result.get("ok") else None
            self.current_width = result.get("width") if result.get("ok") else None
            self.current_height = result.get("height") if result.get("ok") else None
            return result

    def _release_locked(self):
        """Lepas backend yang sedang aktif (kalau ada). Harus dipanggil
        dalam keadaan self.lock sudah dipegang."""
        if self.cap is not None:
            self.cap.release()
            self.cap = None
        if self._tc_handle is not None:
            try:
                self._tc_handle.Close()
            except Exception:
                pass
            self._tc_handle = None
        self._tc_connected = False
        self.backend = None

    def _open_opencv_locked(self, cv_index, width, height):
        backend = cv2.CAP_DSHOW if os.name == "nt" else 0
        cap = cv2.VideoCapture(cv_index, backend)
        if not cap.isOpened():
            return {"ok": False}

        cap.set(cv2.CAP_PROP_FOURCC, cv2.VideoWriter_fourcc(*"MJPG"))

        if width and height:
            cap.set(cv2.CAP_PROP_FRAME_WIDTH, int(width))
            cap.set(cv2.CAP_PROP_FRAME_HEIGHT, int(height))
        actual_w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
        actual_h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        self.cap = cap
        self.backend = "opencv"
        self._opencv_connected = True
        self._opencv_read_failures = 0
        return {"ok": True, "width": actual_w, "height": actual_h}

    def _open_toupcam_locked(self, cam_id, width, height):
        sdk = _load_toupcam_sdk()
        if sdk is None:
            return {"ok": False, "message": "ToupCam SDK tidak tersedia."}
        try:
            h = sdk.Toupcam.Open(cam_id)
        except Exception as e:
            return {"ok": False, "message": f"Gagal membuka kamera (ToupCam): {e}"}
        if h is None:
            return {"ok": False}

        try:
            # Paksa BGR (default di Linux adalah RGB) biar konsisten dengan
            # OpenCV di sisa aplikasi (imencode/imwrite/cvtColor semua asumsi BGR).
            h.put_Option(sdk.TOUPCAM_OPTION_BYTEORDER, 1)
        except Exception:
            pass

        if width and height:
            try:
                # Sensor ToupCam cuma mendukung daftar resolusi diskrit
                # tertentu (bukan sembarang angka) -- put_Size(width, height)
                # dengan nilai yang bukan salah satu step itu akan ditolak
                # (E_INVALIDARG). Jadi cari dulu resolusi valid TERDEKAT dari
                # daftar yang benar-benar didukung kamera, baru pakai
                # put_eSize(index) buat pilih itu.
                target_w, target_h = int(width), int(height)
                n = h.ResolutionNumber()
                best_idx, best_score = None, None
                for i in range(n):
                    rw, rh = h.get_Resolution(i)
                    score = abs(rw - target_w) + abs(rh - target_h)
                    if best_score is None or score < best_score:
                        best_idx, best_score = i, score
                if best_idx is not None:
                    h.put_eSize(best_idx)
            except Exception as e:
                print(f"[DEBUG] ToupCam set resolusi gagal (lanjut pakai resolusi default): {e}")

        try:
            # RealTime = SDK selalu kirim frame TERBARU dan buang backlog,
            h.put_RealTime(1)
        except Exception:
            pass

        try:
            w, ht = h.get_Size()
        except Exception as e:
            try:
                h.Close()
            except Exception:
                pass
            return {"ok": False, "message": f"Gagal membaca ukuran gambar (ToupCam): {e}"}

        self._tc_handle = h
        self._tc_row_pitch = w * 3
        self._tc_buf = bytes(self._tc_row_pitch * ht)
        self._tc_width = w
        self._tc_height = ht
        self._tc_connected = True
        self._tc_last_pull = 0.0
        self.backend = "toupcam"

        try:
            h.StartPullModeWithCallback(CameraManager._toupcam_event_callback, self)
        except Exception as e:
            self._tc_handle = None
            self._tc_connected = False
            self.backend = None
            try:
                h.Close()
            except Exception:
                pass
            return {"ok": False, "message": f"Gagal memulai stream (ToupCam): {e}"}

        return {"ok": True, "width": w, "height": ht}

    @staticmethod
    def _toupcam_event_callback(nEvent, ctx):
        # Callback ini dipanggil dari thread internal SDK, bukan thread Flask.
        ctx._on_toupcam_event(nEvent)

    def _on_toupcam_event(self, nEvent):
        sdk = _load_toupcam_sdk()
        if sdk is None:
            return
        if nEvent == sdk.TOUPCAM_EVENT_IMAGE:
            # Throttle ke ~20fps.
            now = time.monotonic()
            if now - self._tc_last_pull < 0.05:
                return
            self._tc_last_pull = now
            try:
                with self.lock:
                    if self._tc_handle is None:
                        return
                    self._tc_handle.PullImageV4(self._tc_buf, 0, 24, self._tc_row_pitch, None)
                    frame = np.frombuffer(self._tc_buf, dtype=np.uint8).reshape(
                        (self._tc_height, self._tc_width, 3)
                    ).copy()
                    self.last_frame = frame
            except Exception as e:
                print(f"[DEBUG] ToupCam PullImageV4 gagal: {e}")
        elif nEvent in (sdk.TOUPCAM_EVENT_DISCONNECTED, sdk.TOUPCAM_EVENT_ERROR):
            print(f"[DEBUG] ToupCam event disconnect/error: {nEvent}")
            with self.lock:
                self._tc_connected = False
                self.last_frame = None

    def read(self):
        with self.lock:
            if self.backend == "toupcam":
                if not self._tc_connected or self._tc_handle is None:
                    return None
                return self.last_frame
            if self.cap is None or not self.cap.isOpened():
                return None
            ok, frame = self.cap.read()
            if not ok:
                self._opencv_read_failures += 1
                if self._opencv_read_failures >= OPENCV_DISCONNECT_THRESHOLD:
                    self._opencv_connected = False
                    self.last_frame = None
                return None
            self._opencv_read_failures = 0
            self._opencv_connected = True
            self.last_frame = frame
            return frame

    def capture(self):
        with self.lock:
            if self.last_frame is None:
                return None
            self.captured_frame = self.last_frame.copy()
            self.captured_focus_score = compute_focus_score(self.captured_frame)
            return self.captured_frame

    def is_connected(self):
        if self.backend == "toupcam":
            return self._tc_handle is not None and self._tc_connected
        return self.cap is not None and self.cap.isOpened() and self._opencv_connected

    def open_settings_dialog(self):
        """Buka dialog properti bawaan driver DirectShow kamera (kalau didukung).

        Berguna buat atur exposure/white balance/dll manual langsung lewat
        kontrol bawaan drivernya sendiri. Cuma jalan di Windows + backend
        CAP_DSHOW; kamera lain (mis. webcam bawaan laptop) biasanya tidak
        butuh ini karena sudah nurut lewat cv2.set() biasa.
        """
        if os.name != "nt":
            return {"ok": False, "message": "Fitur ini hanya didukung di Windows (DirectShow)."}
        with self.lock:
            if self.cap is None or not self.cap.isOpened():
                return {"ok": False, "message": "Kamera belum terhubung."}
            try:
                self.cap.set(cv2.CAP_PROP_SETTINGS, 1)
            except Exception as e:
                return {"ok": False, "message": f"Gagal membuka dialog pengaturan: {e}"}
            actual_w = int(self.cap.get(cv2.CAP_PROP_FRAME_WIDTH))
            actual_h = int(self.cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
            return {"ok": True, "width": actual_w, "height": actual_h}


camera = CameraManager()


def next_iteration_from_folder(folder, slug):
    """Tentukan nomor iterasi berikutnya dengan membaca file yang SUDAH ADA di folder tujuan.
    """
    pattern = re.compile(
        rf"^{re.escape(slug)}_\d{{8}}-\d{{6}}_(\d+)\.png$", re.IGNORECASE
    )
    max_iter = 0
    try:
        for fname in os.listdir(folder):
            m = pattern.match(fname)
            if m:
                max_iter = max(max_iter, int(m.group(1)))
    except OSError:
        pass
    return max_iter + 1


def gen_mjpeg():
    while True:
        frame = camera.read()
        if frame is None:
            time.sleep(0.2)
            continue
        ok, buf = cv2.imencode(".jpg", frame)
        if not ok:
            continue
        yield (
            b"--frame\r\n"
            b"Content-Type: image/jpeg\r\n\r\n" + buf.tobytes() + b"\r\n"
        )
        time.sleep(0.03)  # kurang lebih 30 fps


def frame_to_png_base64(frame):
    ok, buf = cv2.imencode(".png", frame)
    if not ok:
        return None
    return base64.b64encode(buf.tobytes()).decode("ascii")


def safe_slug(name):
    name = name.strip().lower().replace(" ", "-")
    name = re.sub(r"[^a-z0-9\-]", "", name)
    return name or "pasien"


def compute_focus_score(frame):
    """Estimasi ketajaman fokus."""
    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    return float(cv2.Laplacian(gray, cv2.CV_64F).var())


def resize_for_save(frame, target_width, target_height):
    """Resize gambar sebelum disimpan."
    """
    if not target_width or not target_height:
        return frame
    h, w = frame.shape[:2]
    target_w, target_h = int(target_width), int(target_height)
    if w <= target_w and h <= target_h:
        return frame
    scale = min(target_w / w, target_h / h)
    new_w = max(1, round(w * scale))
    new_h = max(1, round(h * scale))
    return cv2.resize(frame, (new_w, new_h), interpolation=cv2.INTER_AREA)


def is_blank_frame(frame):
    """Deteksi kasar frame yang kemungkinan kosong/lensa tertutup"""
    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    return float(gray.std()) < BLANK_STD_THRESHOLD


def log_capture_csv(folder, row):
    """Tambahkan satu baris metadata ke capture_log.csv di folder tujuan."""
    log_path = os.path.join(folder, LOG_FILENAME)
    is_new = not os.path.exists(log_path)
    fieldnames = [
        "timestamp", "patient_id", "filename", "iterasi",
        "camera", "resolution", "focus_score",
    ]
    with open(log_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        if is_new:
            writer.writeheader()
        writer.writerow(row)


# ---------------------------------------------------------------------------
# Routes - halaman
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    # Dialog "Pengaturan Driver" cuma didukung Windows+DirectShow -- dipakai
    # template buat sembunyikan tombolnya di OS lain (Linux/toupcam) daripada
    # nampilin tombol yang kalau diklik cuma keluar toast "tidak didukung".
    return render_template("index.html", is_windows=(os.name == "nt"))


# ---------------------------------------------------------------------------
# Routes - API kamera
# ---------------------------------------------------------------------------
@app.route("/api/cameras")
def api_cameras():
    return jsonify(camera.list_cameras())


@app.route("/api/camera/connect", methods=["POST"])
def api_camera_connect():
    data = request.json or {}
    index = int(data.get("index", 0))
    width = data.get("width")
    height = data.get("height")
    result = camera.open(index, width, height)
    return jsonify({
        "connected": result["ok"],
        "index": index,
        "width": result.get("width"),
        "height": result.get("height"),
        "message": result.get("message"),
    })


@app.route("/api/camera/settings", methods=["POST"])
def api_camera_settings():
    """Buka dialog properti native driver kamera (lihat CameraManager.open_settings_dialog).
    Dialog ini modal -- request bakal nge-hold sampai user menutup jendelanya."""
    return jsonify(camera.open_settings_dialog())


@app.route("/api/camera/status")
def api_camera_status():
    connected = camera.is_connected()
    return jsonify({
        "connected": connected,
        "index": camera.current_index if connected else None,
        "width": camera.current_width if connected else None,
        "height": camera.current_height if connected else None,
    })


@app.route("/stream")
def stream():
    return Response(gen_mjpeg(), mimetype="multipart/x-mixed-replace; boundary=frame")


# ---------------------------------------------------------------------------
# Routes - capture & save
# ---------------------------------------------------------------------------
@app.route("/api/capture", methods=["POST"])
def api_capture():
    frame = camera.capture()
    if frame is None:
        return jsonify({"ok": False, "message": "Belum ada frame dari kamera."}), 400
    b64 = frame_to_png_base64(frame)
    focus_score = camera.captured_focus_score
    return jsonify({
        "ok": True,
        "image": b64,
        "focusScore": round(focus_score, 1),
        "focusLabel": "Baik" if focus_score >= FOCUS_BLUR_THRESHOLD else "Buram",
        "blankWarning": is_blank_frame(frame),
    })


@app.route("/api/select-folder", methods=["POST"])
def api_select_folder():
    """Buka dialog folder NATIVE lewat tkinter."""
    result = {"path": None}

    def open_dialog():
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        result["path"] = filedialog.askdirectory(title="Pilih Lokasi Penyimpanan")
        root.destroy()

    t = threading.Thread(target=open_dialog)
    t.start()
    t.join()

    if not result["path"]:
        return jsonify({"ok": False}), 400
    return jsonify({"ok": True, "path": result["path"]})


@app.route("/api/save", methods=["POST"])
def api_save():
    data = request.json or {}
    patient_name = (data.get("patientName") or "").strip()
    folder = (data.get("folder") or "").strip()
    camera_name = (data.get("cameraName") or "").strip()
    target_width = data.get("width")
    target_height = data.get("height")

    if not patient_name:
        return jsonify({"ok": False, "message": "Nama pasien belum diisi."}), 400
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Lokasi penyimpanan belum valid."}), 400
    if camera.captured_frame is None:
        return jsonify({"ok": False, "message": "Belum ada gambar yang di-capture."}), 400

    slug = safe_slug(patient_name)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")

    # Semua langkah di bawah ini (hitung iterasi berikutnya dari isi folder,
    # tulis file gambar, tulis baris log CSV) harus jadi satu unit atomik --
    # kalau dua request /api/save nyaris bersamaan (Flask threaded=True)
    # jalan tanpa lock ini, keduanya bisa saja baca "iterasi berikutnya" yang
    # sama sebelum salah satu selesai menulis, jadi filename ketimpa/duplikat
    # atau baris CSV keselip/korup.
    with save_lock:
        iterasi = next_iteration_from_folder(folder, slug)

        # Kalau kamera tidak bisa diatur resolusinya langsung (mis. MiiCam),
        # ini fallback: kecilin gambar secara software persis sebelum disimpan.
        frame_to_save = resize_for_save(camera.captured_frame, target_width, target_height)

        filename = f"{slug}_{timestamp}_{iterasi:02d}.png"
        full_path = os.path.join(folder, filename)

        ok = cv2.imwrite(full_path, frame_to_save)
        if not ok:
            return jsonify({"ok": False, "message": "Gagal menyimpan file."}), 500

        height, width = frame_to_save.shape[:2]
        log_capture_csv(folder, {
            "timestamp": timestamp,
            "patient_id": slug,
            "filename": filename,
            "iterasi": iterasi,
            "camera": camera_name or "?",
            "resolution": f"{width}x{height}",
            "focus_score": round(camera.captured_focus_score or 0, 1),
        })

    return jsonify(
        {
            "ok": True,
            "filename": filename,
            "path": full_path,
            "count": iterasi,
            "width": width,
            "height": height,
        }
    )


@app.route("/api/folder-summary")
def api_folder_summary():
    """Ringkasan jumlah gambar per pasien di folder tujuan, dibaca langsung
    dari nama file yang ada (bukan dari memori/log), jadi selalu akurat."""
    folder = (request.args.get("folder") or "").strip()
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Folder tidak valid."}), 400

    pattern = re.compile(r"^([a-z0-9\-]+)_\d{8}-\d{6}_(\d+)\.png$", re.IGNORECASE)
    counts = {}
    for fname in os.listdir(folder):
        m = pattern.match(fname)
        if m:
            slug = m.group(1)
            counts[slug] = counts.get(slug, 0) + 1

    summary = [{"id": k, "count": v} for k, v in sorted(counts.items())]
    return jsonify({
        "ok": True,
        "summary": summary,
        "totalPatients": len(summary),
        "totalFiles": sum(counts.values()),
    })


def generate_mock_instance_segmentation(frame):
    """PLACEHOLDER -- BUKAN model beneran.

    Rencananya di sini nanti dipasang inference Mask R-CNN/Detectron2 yang
    dilatih dari RBC Chula dataset. Buat sekarang, fungsi ini menghasilkan
    deteksi instance segmentation dummy (blob mirip lingkaran + kelas acak)
    supaya seluruh alur UI (overlay per kelas, daftar per-sel, koreksi
    manual, simpan hasil) bisa dibangun & dites end-to-end lebih dulu.

    Kontrak return-nya (list of dict dengan key classIndex/classLabel/
    confidence/bbox/mask) dipertahankan supaya nanti tinggal ganti ISI
    fungsi ini jadi manggil inference model beneran, tanpa perlu ubah kode
    di route atau di frontend.
    """
    h, w = frame.shape[:2]
    n_cells = random.randint(12, 24)
    detections = []

    # Bobot supaya "Normal cell" (index 0) lebih sering muncul, kira-kira
    # mirip proporsi asli di citra apusan darah (mayoritas sel normal).
    weights = [40] + [5] * (len(RBC_CLASS_LABELS) - 1)

    min_r = max(6, int(min(w, h) * 0.02))
    max_r = max(min_r + 4, int(min(w, h) * 0.045))

    for i in range(n_cells):
        cx = random.randint(max_r, max(max_r + 1, w - max_r))
        cy = random.randint(max_r, max(max_r + 1, h - max_r))
        r = random.randint(min_r, max_r)
        class_index = random.choices(range(len(RBC_CLASS_LABELS)), weights=weights, k=1)[0]
        confidence = round(random.uniform(0.55, 0.98), 3)

        # Bentuk mask sebagai polygon (bukan lingkaran sempurna, dikasih
        # noise dikit di radius) -- ini format yang sama dipakai instance
        # segmentation beneran (polygon per instance, gaya COCO).
        n_points = 16
        points = []
        for p in range(n_points):
            angle = 2 * math.pi * p / n_points
            rr = r * random.uniform(0.85, 1.1)
            points.append([
                round(cx + rr * math.cos(angle), 1),
                round(cy + rr * math.sin(angle), 1),
            ])

        bbox_x = max(0, cx - r)
        bbox_y = max(0, cy - r)
        detections.append({
            "id": i,
            "classIndex": class_index,
            "classLabel": RBC_CLASS_LABELS[class_index],
            "confidence": confidence,
            "bbox": [bbox_x, bbox_y, min(2 * r, w - bbox_x), min(2 * r, h - bbox_y)],
            "mask": points,
        })
    return detections


# ---------------------------------------------------------------------------
# Routes - segmentasi & klasifikasi morfologi sel
# ---------------------------------------------------------------------------
# Scaffold UI + hasil DUMMY dulu -- model instance segmentation beneran
# (rencana: Mask R-CNN/Detectron2, dilatih dari RBC Chula dataset) belum
# dilatih/dipasang. Kontrak response endpoint-endpoint ini sudah dirancang
# final (bukan sekadar placeholder acak) supaya begitu model beneran
# dipasang, frontend nggak perlu diubah lagi.
@app.route("/api/segmentation/classes")
def api_segmentation_classes():
    return jsonify([
        {"index": i, "label": label, "color": RBC_CLASS_COLORS[i]}
        for i, label in enumerate(RBC_CLASS_LABELS)
    ])


@app.route("/api/segmentation/select-image", methods=["POST"])
def api_segmentation_select_image():
    """Buka dialog pilih file gambar NATIVE lewat tkinter, baca isinya, dan
    kirim balik sebagai base64 buat preview (konsisten dengan pola
    /api/capture yang sudah ada -- bukan serve file lewat URL)."""
    result = {"path": None}

    def open_dialog():
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        result["path"] = filedialog.askopenfilename(
            title="Pilih Citra Apusan Darah",
            filetypes=[
                ("Gambar", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff"),
                ("Semua file", "*.*"),
            ],
        )
        root.destroy()

    t = threading.Thread(target=open_dialog)
    t.start()
    t.join()

    if not result["path"]:
        return jsonify({"ok": False}), 400

    frame = cv2.imread(result["path"])
    if frame is None:
        return jsonify({"ok": False, "message": "Gagal membaca file gambar."}), 400

    b64 = frame_to_png_base64(frame)
    height, width = frame.shape[:2]
    return jsonify({
        "ok": True,
        "path": result["path"],
        "filename": os.path.basename(result["path"]),
        "image": b64,
        "width": width,
        "height": height,
    })


@app.route("/api/segmentation/run", methods=["POST"])
def api_segmentation_run():
    """PLACEHOLDER -- model instance segmentation beneran (Mask R-CNN/
    Detectron2, dilatih dari RBC Chula dataset) belum dilatih & dipasang.
    Endpoint ini mengembalikan deteksi DUMMY (lihat
    generate_mock_instance_segmentation) supaya seluruh alur UI bisa
    dibangun & dites duluan. Begitu modelnya siap, ganti isi fungsi ini jadi
    manggil inference beneran -- kontrak response-nya (detections +
    classCounts) dipertahankan supaya frontend nggak perlu diubah lagi."""
    data = request.json or {}
    path = (data.get("path") or "").strip()
    if not path or not os.path.isfile(path):
        return jsonify({"ok": False, "message": "Gambar sumber tidak valid."}), 400

    frame = cv2.imread(path)
    if frame is None:
        return jsonify({"ok": False, "message": "Gagal membaca gambar sumber."}), 400

    # Simulasi waktu proses model beneran, biar indikator progres di UI ada
    # gunanya buat dites -- kecilkan/hapus begitu inference asli dipasang.
    time.sleep(1.2)

    detections = generate_mock_instance_segmentation(frame)
    class_counts = {label: 0 for label in RBC_CLASS_LABELS}
    for d in detections:
        class_counts[d["classLabel"]] += 1

    height, width = frame.shape[:2]
    return jsonify({
        "ok": True,
        "mock": True,
        "message": (
            "Hasil DUMMY/Placeholder"
        ),
        "imageWidth": width,
        "imageHeight": height,
        "detections": detections,
        "classCounts": class_counts,
    })


def _load_segmentation_detail_file(detail_path):
    """Logic bersama buat baca satu file JSON hasil segmentasi + coba baca
    ulang gambar sumbernya buat preview. Dipakai baik dari dialog
    (/api/segmentation/load) maupun klik langsung di daftar hasil folder
    (/api/segmentation/load-path, tanpa buka dialog lagi -- ini yang bikin
    review banyak hasil sekaligus, mis. ~20 capture per pasien, nggak perlu
    buka-tutup dialog file satu-satu)."""
    try:
        with open(detail_path, "r", encoding="utf-8") as f:
            detail = json.load(f)
    except (OSError, json.JSONDecodeError) as e:
        return {"ok": False, "message": f"Gagal membaca file: {e}"}

    source_image = detail.get("sourceImage")
    detections = detail.get("detections") or []

    response = {
        "ok": True,
        "path": source_image,
        "detailFile": os.path.basename(detail_path),
        "patientId": detail.get("patientId"),
        "detections": detections,
        "image": None,
        "width": None,
        "height": None,
    }

    if source_image and os.path.isfile(source_image):
        frame = cv2.imread(source_image)
        if frame is not None:
            response["image"] = frame_to_png_base64(frame)
            h, w = frame.shape[:2]
            response["width"] = w
            response["height"] = h

    if response["image"] is None:
        response["message"] = (
            "Gambar sumber asli tidak ditemukan/gagal dibaca di lokasi "
            "semula -- cuma daftar sel yang berhasil dimuat."
        )

    return response


@app.route("/api/segmentation/load", methods=["POST"])
def api_segmentation_load():
    """Buka file JSON hasil segmentasi yang sudah pernah disimpan lewat
    /api/segmentation/save, muat balik detections-nya (termasuk koreksi
    manual sebelumnya kalau ada), dan coba baca ulang gambar sumbernya buat
    preview -- biar user bisa lihat & lanjut koreksi tanpa mulai dari nol."""
    result = {"path": None}

    def open_dialog():
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        result["path"] = filedialog.askopenfilename(
            title="Pilih File Hasil Segmentasi (JSON)",
            filetypes=[("Hasil segmentasi (JSON)", "*.json"), ("Semua file", "*.*")],
        )
        root.destroy()

    t = threading.Thread(target=open_dialog)
    t.start()
    t.join()

    if not result["path"]:
        return jsonify({"ok": False}), 400

    return jsonify(_load_segmentation_detail_file(result["path"]))


@app.route("/api/segmentation/load-path", methods=["POST"])
def api_segmentation_load_path():
    """Sama seperti /api/segmentation/load, tapi path file JSON-nya dikirim
    langsung dari frontend (klik salah satu baris di daftar "Hasil di
    Folder Ini"), TANPA buka dialog file lagi -- dipakai buat review hasil
    batch (banyak gambar sekaligus, mis. satu pasien ~20 capture) tanpa
    perlu buka-tutup dialog satu-satu per gambar."""
    data = request.json or {}
    detail_path = (data.get("path") or "").strip()
    if not detail_path or not os.path.isfile(detail_path):
        return jsonify({"ok": False, "message": "File hasil tidak valid."}), 400
    return jsonify(_load_segmentation_detail_file(detail_path))


@app.route("/api/segmentation/list-results")
def api_segmentation_list_results():
    """Daftar hasil segmentasi TERBARU per gambar sumber di satu folder
    (dibaca dari segmentation_log.csv) -- ditampilkan sebagai daftar yang
    bisa diklik satu-satu di dalam aplikasi (lewat /load-path).

    Satu gambar bisa saja tercatat berkali-kali di log (mis. diproses
    individual dulu, terus ikut ke-proses lagi lewat batch, atau batch
    dijalankan dua kali) -- karena hasilnya masih DUMMY/acak, tiap proses
    ulang menghasilkan deteksi yang beda. Supaya nggak ambigu "yang mana
    yang berlaku", di sini CUMA baris dengan timestamp TERBARU per
    source_image yang ditampilkan; baris-baris lama tetap ada di CSV
    (nggak dihapus, buat jejak audit) tapi disembunyikan dari daftar ini."""
    folder = (request.args.get("folder") or "").strip()
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Folder tidak valid."}), 400

    try:
        rows = _latest_segmentation_rows(folder)
    except OSError as e:
        return jsonify({"ok": False, "message": f"Gagal membaca log: {e}"}), 500

    results = [
        {
            "detailPath": os.path.join(folder, row.get("detail_file")),
            "detailFile": row.get("detail_file"),
            "sourceImage": row.get("source_image"),
            "patientId": row.get("patient_id"),
            "totalCells": row.get("total_cells"),
            "timestamp": row.get("timestamp"),
            "isDuplicate": (row.get("is_duplicate") or "0") == "1",
            "duplicateOfFile": row.get("duplicate_of") or None,
        }
        for row in rows
        if row.get("detail_file")
    ]
    results.sort(key=lambda r: r["timestamp"] or "", reverse=True)
    return jsonify({"ok": True, "folder": folder, "results": results})


@app.route("/api/segmentation/save", methods=["POST"])
def api_segmentation_save():
    """Simpan hasil segmentasi (termasuk koreksi manual dari UI, kalau ada)
    ke dua tempat, keduanya di folder yang sama dengan gambar sumbernya:
      1. File JSON detail per-sel (mask polygon, confidence, status koreksi)
      2. Satu baris ringkasan di segmentation_log.csv (jumlah sel per kelas
         -- format kolom per kelas biar gampang dianalisis langsung di Excel)

    Sebelum benar-benar nulis, dicek dulu apakah hasil ini PERSIS SAMA
    (total sel + jumlah per kelas identik) dengan hasil TERAKHIR yang sudah
    tersimpan buat gambar yang sama -- kalau iya dan client belum kirim
    `force: true`, nggak jadi ditulis, cuma balikin `needsConfirmation` biar
    frontend nanya dulu ke user (lihat doSaveSegmentation()). Ini buat
    nyegah folder numpuk banyak file JSON isinya sama persis tanpa sadar.
    """
    data = request.json or {}
    path = (data.get("path") or "").strip()
    patient_id = (data.get("patientId") or "").strip()
    detections = data.get("detections") or []
    force = bool(data.get("force"))

    if not path or not os.path.isfile(path):
        return jsonify({"ok": False, "message": "Gambar sumber tidak valid."}), 400
    if not detections:
        return jsonify({"ok": False, "message": "Belum ada hasil segmentasi untuk disimpan."}), 400

    folder = os.path.dirname(path)
    base_name = os.path.splitext(os.path.basename(path))[0]
    source_image_basename = os.path.basename(path)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")

    class_counts = {label: 0 for label in RBC_CLASS_LABELS}
    for d in detections:
        label = d.get("correctedClassLabel") or d.get("classLabel")
        if label in class_counts:
            class_counts[label] += 1

    # Agregasi penilaian dokter per sel (doctorVerdict: "correct"/"incorrect"/
    # null kalau belum direview) -- dipakai buat kolom ringkasan di
    # segmentation_log.csv & laporan Excel (lihat api_segmentation_export_report).
    doctor_correct_count = sum(1 for d in detections if d.get("doctorVerdict") == "correct")
    doctor_incorrect_count = sum(1 for d in detections if d.get("doctorVerdict") == "incorrect")
    doctor_reviewed_count = doctor_correct_count + doctor_incorrect_count

    with segmentation_log_lock:
        log_path = os.path.join(folder, SEGMENTATION_LOG_FILENAME)
        _ensure_log_columns(log_path, SEGMENTATION_LOG_FIELDNAMES)

        prev_row = next(
            (r for r in _latest_segmentation_rows(folder) if r.get("source_image") == source_image_basename),
            None,
        )
        is_duplicate = False
        duplicate_of = ""
        if prev_row is not None:
            try:
                prev_total = int(prev_row.get("total_cells") or 0)
                counts_match = all(
                    int(prev_row.get(label) or 0) == class_counts[label] for label in RBC_CLASS_LABELS
                )
                is_duplicate = counts_match and prev_total == len(detections)
            except (TypeError, ValueError):
                is_duplicate = False
            if is_duplicate:
                duplicate_of = prev_row.get("detail_file") or ""

        if is_duplicate and not force:
            return jsonify({
                "ok": True,
                "needsConfirmation": True,
                "isDuplicate": True,
                "lastTimestamp": prev_row.get("timestamp"),
                "lastDetailFile": duplicate_of,
            })

        detail_filename = f"{base_name}_segmentation_{timestamp}.json"
        detail_path = os.path.join(folder, detail_filename)
        # Timestamp cuma presisi per detik -- kalau ada dua save dalam detik
        # yang sama (mis. warning duplikat lalu langsung "tetap simpan"),
        # nama filenya bakal sama & nimpa file sebelumnya. Tambah suffix
        # numerik biar tetap unik.
        dedup_suffix = 1
        while os.path.exists(detail_path):
            detail_filename = f"{base_name}_segmentation_{timestamp}-{dedup_suffix}.json"
            detail_path = os.path.join(folder, detail_filename)
            dedup_suffix += 1
        try:
            with open(detail_path, "w", encoding="utf-8") as f:
                json.dump({
                    "timestamp": timestamp,
                    "patientId": patient_id,
                    "sourceImage": path,
                    "detections": detections,
                    "isDuplicate": is_duplicate,
                    "duplicateOfFile": duplicate_of or None,
                }, f, ensure_ascii=False, indent=2)
        except OSError as e:
            return jsonify({"ok": False, "message": f"Gagal menyimpan file detail: {e}"}), 500

        is_new = not os.path.exists(log_path)
        row = {
            "timestamp": timestamp,
            "patient_id": patient_id or "?",
            "source_image": source_image_basename,
            "detail_file": detail_filename,
            "total_cells": len(detections),
            "is_duplicate": "1" if is_duplicate else "0",
            "duplicate_of": duplicate_of,
            "doctor_reviewed_count": doctor_reviewed_count,
            "doctor_correct_count": doctor_correct_count,
            "doctor_incorrect_count": doctor_incorrect_count,
            **class_counts,
        }
        try:
            with open(log_path, "a", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=SEGMENTATION_LOG_FIELDNAMES)
                if is_new:
                    writer.writeheader()
                writer.writerow(row)
        except OSError as e:
            return jsonify({"ok": False, "message": f"Gagal menulis log: {e}"}), 500

    return jsonify({
        "ok": True,
        "detailPath": detail_path,
        "logPath": log_path,
        "totalCells": len(detections),
        "isDuplicate": is_duplicate,
    })


@app.route("/api/segmentation/batch-select", methods=["POST"])
def api_segmentation_batch_select():
    """Dialog pilih FOLDER buat proses batch -- daftar file gambar langsung
    di dalam folder itu (bukan rekursif ke subfolder) dikembalikan supaya
    frontend bisa konfirmasi dulu jumlahnya ke user sebelum benar-benar
    diproses (bisa makan waktu kalau isinya banyak)."""
    result = {"path": None}

    def open_dialog():
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        result["path"] = filedialog.askdirectory(title="Pilih Folder Citra untuk Diproses Batch")
        root.destroy()

    t = threading.Thread(target=open_dialog)
    t.start()
    t.join()

    if not result["path"]:
        return jsonify({"ok": False}), 400

    image_exts = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff")
    try:
        files = sorted(
            f for f in os.listdir(result["path"])
            if f.lower().endswith(image_exts) and os.path.isfile(os.path.join(result["path"], f))
        )
    except OSError as e:
        return jsonify({"ok": False, "message": f"Gagal membaca isi folder: {e}"}), 400

    return jsonify({"ok": True, "folder": result["path"], "files": files})


@app.route("/api/segmentation/batch-run", methods=["POST"])
def api_segmentation_batch_run():
    """PLACEHOLDER -- proses SEMUA gambar dalam satu folder sekaligus:
    jalankan (mock) instance segmentation + langsung simpan hasilnya (JSON
    detail + baris di segmentation_log.csv) untuk tiap gambar, tanpa jeda
    buat koreksi manual per gambar (koreksi tetap bisa dilakukan belakangan
    per file lewat "Muat Hasil Tersimpan"). Model instance segmentation
    beneran belum dilatih -- begitu siap, ganti pemanggilan
    generate_mock_instance_segmentation() di bawah jadi inference asli,
    sisanya (penyimpanan, agregasi) nggak perlu diubah."""
    data = request.json or {}
    folder = (data.get("folder") or "").strip()
    patient_id = (data.get("patientId") or "").strip()
    skip_existing = bool(data.get("skipExisting"))
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Folder tidak valid."}), 400

    image_exts = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff")
    files = sorted(
        f for f in os.listdir(folder)
        if f.lower().endswith(image_exts) and os.path.isfile(os.path.join(folder, f))
    )
    if not files:
        return jsonify({"ok": False, "message": "Tidak ada file gambar di folder ini."}), 400

    grand_total_counts = {label: 0 for label in RBC_CLASS_LABELS}
    per_image_results = []
    already_processed = set()

    with segmentation_log_lock:
        log_path = os.path.join(folder, SEGMENTATION_LOG_FILENAME)
        _ensure_log_columns(log_path, SEGMENTATION_LOG_FIELDNAMES)
        is_new_log = not os.path.exists(log_path)

        if skip_existing and os.path.isfile(log_path):
            try:
                with open(log_path, "r", newline="", encoding="utf-8") as f:
                    for row in csv.DictReader(f):
                        si = row.get("source_image")
                        if si:
                            already_processed.add(si)
            except OSError:
                pass

        for fname in files:
            if skip_existing and fname in already_processed:
                per_image_results.append({
                    "filename": fname, "ok": True, "skipped": True,
                    "message": "Dilewati -- sudah pernah diproses sebelumnya.",
                })
                continue

            fpath = os.path.join(folder, fname)
            frame = cv2.imread(fpath)
            if frame is None:
                per_image_results.append({"filename": fname, "ok": False, "message": "Gagal dibaca."})
                continue

            # Delay simulasi diperkecil dibanding mode satu-gambar, biar
            # batch banyak file nggak kelamaan -- hapus kalau model beneran
            # sudah dipasang (delay ini murni buat placeholder).
            time.sleep(0.3)
            detections = generate_mock_instance_segmentation(frame)

            base_name = os.path.splitext(fname)[0]
            timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
            detail_filename = f"{base_name}_segmentation_{timestamp}.json"
            detail_path = os.path.join(folder, detail_filename)
            try:
                with open(detail_path, "w", encoding="utf-8") as f:
                    json.dump({
                        "timestamp": timestamp,
                        "patientId": patient_id,
                        "sourceImage": fpath,
                        "detections": detections,
                    }, f, ensure_ascii=False, indent=2)
            except OSError as e:
                per_image_results.append({"filename": fname, "ok": False, "message": f"Gagal menyimpan detail: {e}"})
                continue

            class_counts = {label: 0 for label in RBC_CLASS_LABELS}
            for d in detections:
                class_counts[d["classLabel"]] += 1
                grand_total_counts[d["classLabel"]] += 1

            row = {
                "timestamp": timestamp,
                "patient_id": patient_id or "?",
                "source_image": fname,
                "detail_file": detail_filename,
                "total_cells": len(detections),
                "is_duplicate": "0",  # batch selalu jalankan segmentasi baru, nggak dicek duplikat
                "duplicate_of": "",
                # Hasil batch murni mock otomatis, belum pernah direview dokter
                # satu-satu -- kolomnya tetap diisi 0 biar CSV konsisten.
                "doctor_reviewed_count": 0,
                "doctor_correct_count": 0,
                "doctor_incorrect_count": 0,
                **class_counts,
            }
            try:
                with open(log_path, "a", newline="", encoding="utf-8") as f:
                    writer = csv.DictWriter(f, fieldnames=SEGMENTATION_LOG_FIELDNAMES)
                    if is_new_log:
                        writer.writeheader()
                        is_new_log = False
                    writer.writerow(row)
            except OSError as e:
                per_image_results.append({"filename": fname, "ok": False, "message": f"Gagal menulis log: {e}"})
                continue

            per_image_results.append({
                "filename": fname,
                "ok": True,
                "totalCells": len(detections),
                "classCounts": class_counts,
                "detailFile": detail_filename,
            })

    total_skipped = sum(1 for r in per_image_results if r.get("skipped"))
    total_processed = sum(1 for r in per_image_results if r["ok"] and not r.get("skipped"))
    return jsonify({
        "ok": True,
        "mock": True,
        "folder": folder,
        "totalImages": len(files),
        "totalImagesOk": sum(1 for r in per_image_results if r["ok"]),
        "totalSkipped": total_skipped,
        "grandTotalCounts": grand_total_counts,
        "perImage": per_image_results,
        "message": (
            f"Batch selesai (hasil DUMMY/placeholder) -- {total_processed} gambar diproses"
            + (f", {total_skipped} dilewati (sudah ada)" if total_skipped else "")
            + ". Model instance segmentation beneran belum dilatih."
        ),
    })


@app.route("/api/segmentation/check-existing")
def api_segmentation_check_existing():
    """Cek apakah satu gambar (by path) sudah pernah punya hasil segmentasi
    tersimpan di segmentation_log.csv folder yang sama -- dipakai buat
    warning sebelum reprocess individual (biar nggak nggak sadar bikin
    hasil baru yang beda random dari yang lama, lihat juga skipExisting di
    /batch-run buat kasus batch)."""
    path = (request.args.get("path") or "").strip()
    if not path:
        return jsonify({"exists": False})

    folder = os.path.dirname(path)
    fname = os.path.basename(path)
    log_path = os.path.join(folder, SEGMENTATION_LOG_FILENAME)
    if not os.path.isfile(log_path):
        return jsonify({"exists": False})

    latest = None
    try:
        with open(log_path, "r", newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if row.get("source_image") != fname:
                    continue
                ts = row.get("timestamp") or ""
                if latest is None or ts >= (latest.get("timestamp") or ""):
                    latest = row
    except OSError:
        return jsonify({"exists": False})

    if latest is None:
        return jsonify({"exists": False})

    return jsonify({
        "exists": True,
        "detailFile": latest.get("detail_file"),
        "totalCells": latest.get("total_cells"),
        "timestamp": latest.get("timestamp"),
    })


@app.route("/api/segmentation/patient-summary")
def api_segmentation_patient_summary():
    """Agregasi segmentation_log.csv per patient_id dalam satu folder --
    total gambar, total sel, dan breakdown per kelas.

    Dedup dulu per source_image (ambil baris TERBARU aja per gambar,
    sama seperti /api/segmentation/list-results) SEBELUM diagregasi --
    kalau tidak, gambar yang kebetulan diproses lebih dari sekali
    (individual lalu ikut batch, atau batch dijalankan dua kali) bakal
    kehitung dobel/lebih di total & breakdown per kelas."""
    folder = (request.args.get("folder") or "").strip()
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Folder tidak valid."}), 400

    try:
        rows = _latest_segmentation_rows(folder)
    except OSError as e:
        return jsonify({"ok": False, "message": f"Gagal membaca log: {e}"}), 500

    patients = {}
    for row in rows:
        pid = row.get("patient_id") or "?"
        if pid not in patients:
            patients[pid] = {
                "patientId": pid,
                "totalImages": 0,
                "totalCells": 0,
                "totalReviewed": 0,
                "totalCorrect": 0,
                "totalIncorrect": 0,
                "classCounts": {label: 0 for label in RBC_CLASS_LABELS},
            }
        patients[pid]["totalImages"] += 1
        try:
            patients[pid]["totalCells"] += int(row.get("total_cells") or 0)
        except ValueError:
            pass
        for key, field in (
            ("totalReviewed", "doctor_reviewed_count"),
            ("totalCorrect", "doctor_correct_count"),
            ("totalIncorrect", "doctor_incorrect_count"),
        ):
            try:
                patients[pid][key] += int(row.get(field) or 0)
            except (ValueError, TypeError):
                pass
        for label in RBC_CLASS_LABELS:
            try:
                patients[pid]["classCounts"][label] += int(row.get(label) or 0)
            except (ValueError, TypeError):
                pass

    return jsonify({
        "ok": True,
        "patients": sorted(patients.values(), key=lambda p: p["patientId"]),
    })


def _latest_segmentation_rows(folder):
    """Baca segmentation_log.csv di satu folder, dedup ambil baris TERBARU
    per source_image -- helper yang sama dipakai /list-results dan
    /patient-summary, dipusatkan di sini juga buat /export-report.

    Kalau dua baris punya timestamp yang PERSIS sama (resolusi timestamp
    cuma per detik, bisa kejadian kalau ada save cepat berturut-turut,
    mis. warning duplikat lalu "tetap simpan"), baris yang lebih BELAKANGAN
    di file (ditulis lebih baru) yang menang -- makanya perbandingannya
    `>` bukan `>=`, biar baris lama nggak "menang" cuma gara-gara duluan
    ditulis di file."""
    log_path = os.path.join(folder, SEGMENTATION_LOG_FILENAME)
    if not os.path.isfile(log_path):
        return []
    latest_by_image = {}
    with open(log_path, "r", newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            si = row.get("source_image")
            if not si:
                continue
            ts = row.get("timestamp") or ""
            existing = latest_by_image.get(si)
            if existing is not None and existing.get("timestamp", "") > ts:
                continue
            latest_by_image[si] = row
    return list(latest_by_image.values())


@app.route("/api/segmentation/export-report", methods=["POST"])
def api_segmentation_export_report():
    """Export laporan ringkasan hasil segmentasi satu folder ke file Excel
    (2 sheet: ringkasan per pasien, dan detail per gambar) -- dedup terbaru
    per gambar, konsisten sama yang ditampilkan di aplikasi (Ringkasan per
    Pasien / Hasil di Folder Ini). Kalau patientId diisi, cuma pasien itu
    yang diekspor; kalau kosong, semua pasien di folder itu."""
    data = request.json or {}
    folder = (data.get("folder") or "").strip()
    patient_id_filter = (data.get("patientId") or "").strip()
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Folder tidak valid."}), 400

    rows = _latest_segmentation_rows(folder)
    if patient_id_filter:
        rows = [r for r in rows if (r.get("patient_id") or "") == patient_id_filter]
    if not rows:
        return jsonify({"ok": False, "message": "Tidak ada hasil yang cocok untuk diekspor."}), 400

    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        return jsonify({
            "ok": False,
            "message": "Library openpyxl belum terinstall di server (pip install openpyxl).",
        }), 500

    wb = openpyxl.Workbook()

    doctor_cols = ["Direview Dokter", "Benar (Dokter)", "Salah (Dokter)"]

    ws1 = wb.active
    ws1.title = "Ringkasan Pasien"
    ws1.append(["ID Pasien", "Jumlah Gambar", "Total Sel"] + doctor_cols + RBC_CLASS_LABELS)
    for cell in ws1[1]:
        cell.font = Font(bold=True)

    patients = {}
    for row in rows:
        pid = row.get("patient_id") or "?"
        if pid not in patients:
            patients[pid] = {
                "images": 0,
                "totalCells": 0,
                "reviewed": 0,
                "correct": 0,
                "incorrect": 0,
                "counts": {l: 0 for l in RBC_CLASS_LABELS},
            }
        patients[pid]["images"] += 1
        try:
            patients[pid]["totalCells"] += int(row.get("total_cells") or 0)
        except ValueError:
            pass
        for key in ("reviewed", "correct", "incorrect"):
            try:
                patients[pid][key] += int(row.get(f"doctor_{key}_count") or 0)
            except (ValueError, TypeError):
                pass
        for label in RBC_CLASS_LABELS:
            try:
                patients[pid]["counts"][label] += int(row.get(label) or 0)
            except (ValueError, TypeError):
                pass
    for pid, agg in sorted(patients.items()):
        ws1.append(
            [pid, agg["images"], agg["totalCells"], agg["reviewed"], agg["correct"], agg["incorrect"]]
            + [agg["counts"][l] for l in RBC_CLASS_LABELS]
        )

    ws2 = wb.create_sheet("Detail per Gambar")
    ws2.append(["Timestamp", "ID Pasien", "Gambar", "Total Sel"] + doctor_cols + RBC_CLASS_LABELS)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for row in sorted(rows, key=lambda r: r.get("timestamp") or ""):
        ws2.append(
            [row.get("timestamp"), row.get("patient_id"), row.get("source_image"), row.get("total_cells")]
            + [row.get("doctor_reviewed_count"), row.get("doctor_correct_count"), row.get("doctor_incorrect_count")]
            + [row.get(l) for l in RBC_CLASS_LABELS]
        )

    # Sheet ketiga: satu baris per SEL terdeteksi (bukan per gambar/pasien) --
    # butuh baca ulang tiap file JSON detail (segmentation_log.csv cuma nyimpen
    # ringkasan per gambar, bukan per sel). File yang hilang/rusak dilewati
    # (baris itu aja) supaya satu file bermasalah nggak menggagalkan seluruh
    # export.
    ws3 = wb.create_sheet("Detail per Sel")
    ws3.append([
        "Timestamp", "ID Pasien", "Gambar", "No. Sel", "Kelas Model",
        "Kelas Dikoreksi", "Confidence (%)", "Manual", "Verdict Dokter",
    ])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for row in sorted(rows, key=lambda r: r.get("timestamp") or ""):
        detail_file = row.get("detail_file")
        if not detail_file:
            continue
        detail_path = os.path.join(folder, detail_file)
        try:
            with open(detail_path, "r", encoding="utf-8") as f:
                detail = json.load(f)
        except (OSError, json.JSONDecodeError):
            continue
        for i, d in enumerate(detail.get("detections") or []):
            corrected_idx = d.get("correctedClassIndex")
            has_correction = corrected_idx is not None and corrected_idx != d.get("classIndex")
            verdict = {"correct": "Benar", "incorrect": "Salah"}.get(d.get("doctorVerdict"), "-")
            confidence = d.get("confidence")
            ws3.append([
                row.get("timestamp"),
                row.get("patient_id"),
                row.get("source_image"),
                i + 1,
                d.get("classLabel"),
                d.get("correctedClassLabel") if has_correction else "",
                round(confidence * 100, 1) if isinstance(confidence, (int, float)) else confidence,
                "Ya" if d.get("manual") else "Tidak",
                verdict,
            ])

    for ws in (ws1, ws2, ws3):
        for col_cells in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(30, max(10, length + 2))

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    safe_patient = safe_slug(patient_id_filter) if patient_id_filter else "semua-pasien"
    # Timestamp cuma presisi per detik -- dua export cepat berturut-turut bisa
    # jatuh di detik yang sama dan nimpa file satu sama lain tanpa peringatan
    # kalau nama filenya sama persis (bug yang sama ditemukan & diperbaiki di
    # api_segmentation_export_coco). Tambah suffix numerik biar tetap unik.
    dedup_suffix = 0
    filename = f"laporan_segmentasi_{safe_patient}_{timestamp}.xlsx"
    export_path = os.path.join(folder, filename)
    while os.path.exists(export_path):
        dedup_suffix += 1
        filename = f"laporan_segmentasi_{safe_patient}_{timestamp}-{dedup_suffix}.xlsx"
        export_path = os.path.join(folder, filename)
    try:
        wb.save(export_path)
    except OSError as e:
        return jsonify({"ok": False, "message": f"Gagal menyimpan file laporan: {e}"}), 500

    return jsonify({"ok": True, "path": export_path, "filename": filename})


def _polygon_area(points):
    """Luas polygon sederhana (shoelace formula) -- dipakai buat field "area"
    di annotation COCO, nggak butuh dependency tambahan apapun."""
    n = len(points)
    if n < 3:
        return 0.0
    area = 0.0
    for i in range(n):
        x1, y1 = points[i]
        x2, y2 = points[(i + 1) % n]
        area += x1 * y2 - x2 * y1
    return abs(area) / 2.0


@app.route("/api/segmentation/export-coco", methods=["POST"])
def api_segmentation_export_coco():
    """Export hasil segmentasi satu folder ke satu file annotations.json
    format COCO instance segmentation (images/annotations/categories) --
    siap dipakai buat training Mask R-CNN/Detectron2 nanti begitu model
    aslinya mulai dikerjakan. Beda dari export-report (Excel, buat dibaca
    manusia), ini format standar yang langsung bisa dimuat library training.

    Kelas yang dipakai per sel selalu kelas EFEKTIF-nya (correctedClassLabel
    kalau ada koreksi, kalau tidak classLabel asli model) -- sama seperti
    logic class_counts di api_segmentation_save.

    onlyReviewed=True: cuma sel dengan doctorVerdict terisi yang masuk. Sel
    "incorrect" TANPA correctedClassIndex (dokter cuma tandai salah, belum
    sempat ganti kelasnya) sengaja DILEWATI biar nggak nyemarin dataset
    dengan label yang sudah diketahui keliru -- bukan disertakan pakai label
    asli yang salah itu."""
    data = request.json or {}
    folder = (data.get("folder") or "").strip()
    patient_id_filter = (data.get("patientId") or "").strip()
    only_reviewed = bool(data.get("onlyReviewed"))
    if not folder or not os.path.isdir(folder):
        return jsonify({"ok": False, "message": "Folder tidak valid."}), 400

    rows = _latest_segmentation_rows(folder)
    if patient_id_filter:
        rows = [r for r in rows if (r.get("patient_id") or "") == patient_id_filter]
    if not rows:
        return jsonify({"ok": False, "message": "Tidak ada hasil yang cocok untuk diekspor."}), 400

    category_by_label = {label: i + 1 for i, label in enumerate(RBC_CLASS_LABELS)}
    categories = [{"id": i + 1, "name": label} for i, label in enumerate(RBC_CLASS_LABELS)]

    images = []
    annotations = []
    skipped_images = 0
    skipped_cells = 0
    next_image_id = 1
    next_ann_id = 1

    for row in sorted(rows, key=lambda r: r.get("timestamp") or ""):
        detail_file = row.get("detail_file")
        if not detail_file:
            continue
        detail_path = os.path.join(folder, detail_file)
        try:
            with open(detail_path, "r", encoding="utf-8") as f:
                detail = json.load(f)
        except (OSError, json.JSONDecodeError):
            skipped_images += 1
            continue

        source_image = detail.get("sourceImage")
        frame = cv2.imread(source_image) if source_image else None
        if frame is None:
            skipped_images += 1
            continue
        height, width = frame.shape[:2]

        image_id = next_image_id
        next_image_id += 1
        images.append({
            "id": image_id,
            "file_name": os.path.basename(source_image),
            "width": width,
            "height": height,
        })

        for d in (detail.get("detections") or []):
            corrected_idx = d.get("correctedClassIndex")
            has_correction = corrected_idx is not None and corrected_idx != d.get("classIndex")
            label = d.get("correctedClassLabel") if has_correction else d.get("classLabel")
            verdict = d.get("doctorVerdict")

            if only_reviewed:
                if verdict not in ("correct", "incorrect"):
                    skipped_cells += 1
                    continue
                if verdict == "incorrect" and not has_correction:
                    skipped_cells += 1
                    continue

            category_id = category_by_label.get(label)
            mask = d.get("mask") or []
            if category_id is None or len(mask) < 3:
                skipped_cells += 1
                continue

            bbox = d.get("bbox")
            if not bbox or len(bbox) != 4:
                xs = [p[0] for p in mask]
                ys = [p[1] for p in mask]
                bbox = [min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)]

            annotations.append({
                "id": next_ann_id,
                "image_id": image_id,
                "category_id": category_id,
                "segmentation": [[coord for point in mask for coord in point]],
                "area": _polygon_area(mask),
                "bbox": bbox,
                "iscrowd": 0,
            })
            next_ann_id += 1

    if not images:
        return jsonify({"ok": False, "message": "Tidak ada gambar yang berhasil dibaca untuk diekspor."}), 400

    model_path = (_load_model_config().get("modelPath") or "").strip()
    coco = {
        "info": {
            "description": (
                "Dataset diekspor dari Thalassemia Capture App. "
                + (
                    f'Dihasilkan model mock (belum ada inference asli; modelPath tersimpan: "{model_path}").'
                    if not model_path
                    else f'modelPath tersimpan: "{model_path}" -- cek apakah inference asli sudah terpasang.'
                )
            ),
            "date_created": datetime.now().isoformat(),
        },
        "images": images,
        "annotations": annotations,
        "categories": categories,
    }

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    safe_patient = safe_slug(patient_id_filter) if patient_id_filter else "semua-pasien"
    # Timestamp cuma presisi per detik -- dua export cepat berturut-turut
    # (mis. toggle checkbox lalu export lagi) bisa jatuh di detik yang sama
    # dan nimpa file satu sama lain tanpa peringatan kalau nama filenya sama
    # persis. Tambah suffix numerik biar tetap unik (pola sama seperti
    # detail_filename di api_segmentation_save).
    dedup_suffix = 0
    filename = f"dataset_coco_{safe_patient}_{timestamp}.json"
    export_path = os.path.join(folder, filename)
    while os.path.exists(export_path):
        dedup_suffix += 1
        filename = f"dataset_coco_{safe_patient}_{timestamp}-{dedup_suffix}.json"
        export_path = os.path.join(folder, filename)
    try:
        with open(export_path, "w", encoding="utf-8") as f:
            json.dump(coco, f, ensure_ascii=False, indent=2)
    except OSError as e:
        return jsonify({"ok": False, "message": f"Gagal menyimpan dataset: {e}"}), 500

    skip_note = ""
    if skipped_images or skipped_cells:
        skip_note = (
            f" ({skipped_images} gambar dilewati (tidak terbaca), "
            f"{skipped_cells} sel dilewati (belum direview/koreksi tidak lengkap))"
        )
    return jsonify({
        "ok": True,
        "path": export_path,
        "filename": filename,
        "message": f"Dataset COCO tersimpan: {len(images)} gambar, {len(annotations)} sel{skip_note}.",
    })


# ---------------------------------------------------------------------------
# Status & konfigurasi model segmentasi
# ---------------------------------------------------------------------------
# Placeholder transparansi -- inference model instance segmentation ASLI
# (Mask R-CNN/Detectron2) belum dipasang sama sekali (lihat
# generate_mock_instance_segmentation). File config ini cuma NYIMPAN path
# checkpoint yang direncanakan dipakai nanti, belum benar-benar memuat/
# menjalankan model apapun -- tujuannya biar plumbing-nya sudah siap begitu
# integrasi model beneran dikerjakan (tinggal baca MODEL_CONFIG_PATH di
# generate_mock_instance_segmentation/api_segmentation_run dan branch ke
# inference asli).
MODEL_CONFIG_FILENAME = "model_config.json"
MODEL_CONFIG_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), MODEL_CONFIG_FILENAME
)

# Ambang batas (persen) confidence sel yang dianggap "rendah" di UI (badge +
# garis putus-putus di overlay) -- disimpan di file config yang sama dengan
# modelPath biar bisa diatur dari panel "Status Model Segmentasi" tanpa perlu
# ubah kode, bukan angka mutlak yang berlaku selamanya.
DEFAULT_LOW_CONFIDENCE_THRESHOLD = 70


def _load_model_config():
    config = {"modelPath": "", "lowConfidenceThreshold": DEFAULT_LOW_CONFIDENCE_THRESHOLD}
    if os.path.isfile(MODEL_CONFIG_PATH):
        try:
            with open(MODEL_CONFIG_PATH, "r", encoding="utf-8") as f:
                loaded = json.load(f)
            if isinstance(loaded, dict):
                config.update(loaded)
        except (OSError, json.JSONDecodeError):
            pass
    try:
        threshold = int(config.get("lowConfidenceThreshold"))
        if not (1 <= threshold <= 100):
            raise ValueError
        config["lowConfidenceThreshold"] = threshold
    except (TypeError, ValueError):
        config["lowConfidenceThreshold"] = DEFAULT_LOW_CONFIDENCE_THRESHOLD
    return config


@app.route("/api/segmentation/model-status")
def api_segmentation_model_status():
    config = _load_model_config()
    model_path = (config.get("modelPath") or "").strip()
    model_file_found = bool(model_path) and os.path.isfile(model_path)
    if model_path:
        detail = (
            f'Path model tersimpan ("{model_path}"), tapi kode inference asli belum '
            "dipasang -- ini baru tempat konfigurasi buat nanti."
        )
    else:
        detail = "Belum ada path model dikonfigurasi."
    return jsonify({
        "ok": True,
        "usingMock": True,
        "modelPath": model_path,
        "modelFileFound": model_file_found,
        "lowConfidenceThreshold": config["lowConfidenceThreshold"],
        "message": "Semua hasil segmentasi saat ini masih DUMMY/placeholder. " + detail,
    })


@app.route("/api/segmentation/model-config", methods=["POST"])
def api_segmentation_model_config():
    data = request.json or {}
    model_path = (data.get("modelPath") or "").strip()
    try:
        threshold = int(data.get("lowConfidenceThreshold"))
        if not (1 <= threshold <= 100):
            raise ValueError
    except (TypeError, ValueError):
        threshold = DEFAULT_LOW_CONFIDENCE_THRESHOLD
    try:
        with open(MODEL_CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(
                {"modelPath": model_path, "lowConfidenceThreshold": threshold},
                f, ensure_ascii=False, indent=2,
            )
    except OSError as e:
        return jsonify({"ok": False, "message": f"Gagal menyimpan konfigurasi: {e}"}), 500
    return jsonify({"ok": True, "modelPath": model_path, "lowConfidenceThreshold": threshold})


if __name__ == "__main__":
    app.run(debug=True, threaded=True)
